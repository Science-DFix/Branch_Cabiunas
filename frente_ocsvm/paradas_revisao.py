#!/usr/bin/env python3
"""Lista das paradas da turbina sem alarme crítico, para a equipe dizer quais foram programadas.

Mesma definição da contagem de 14/10/2026: parada = passagem do estado operacional para "off_longo" com a máquina
"on" nas 3 h anteriores (39 no período). As 8 com alarme HIGH/LOLO do catálogo em ±15 min são os 8 trips curados;
as outras 31 vão para revisão. Os dados e o catálogo estão no mesmo relógio (UTC); a planilha mostra também
Brasília (UTC-3).

Saídas em frente_ocsvm/revisao/:
  paradas_para_revisao.xlsx   para a equipe (revisão CEGA: sem a coluna da OCSVM)
  paradas_interno.csv         as mesmas 39, com a coluna "OCSVM alertou nas 48 h antes" (só para nós)

Uso:  python frente_ocsvm/paradas_revisao.py
"""
from __future__ import annotations
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import comum as C

SAIDA = C.AQUI / "revisao"
OPCOES = ["Programada", "Operacional (não programada, sem falha)", "Falha de equipamento", "Não sei"]
FONTE = Font(name="Arial", size=10)
NEGRITO = Font(name="Arial", size=10, bold=True)
AMARELO = PatternFill("solid", fgColor="FFFF00")


def paradas() -> pd.DataFrame:
    df = C.carrega_canais()
    idx, op = df.index, df["operational_state"]
    ft = C.trips(idx)
    on = op == "on"
    a = pd.read_csv(C.DADOS / "alarmes_selecionados_turbina_a.csv")
    a["t"] = pd.to_datetime(a["Data da Ocorrência"], errors="coerce")
    a = a[a.Status.astype(str).str.startswith("ACT") & a.t.between(idx.min(), idx.max())]
    can = {c: df[c].astype(bool) for c in C.CANAIS}
    fin = C.decide(can, **{k: C.REF[k] for k in ("min_votes", "dur_min", "refrat_h")})
    eps = C.group_alerts_into_episodes(fin.astype(int), merge_gap_minutes=120.0)

    mud = op[op.ne(op.shift())]
    L = []
    for t in [t for t, v in mud.items() if v == "off_longo"]:
        if not on[(idx >= t - pd.Timedelta(hours=3)) & (idx < t)].any():
            continue
        volta = on[(idx > t) & on]
        fim = volta.index[0] if len(volta) else idx[-1]
        partida = op[(idx < t) & (op == "off_longo")]  # fim da última parada longa
        x = a[(a.t >= t - pd.Timedelta(minutes=15)) & (a.t <= t + pd.Timedelta(minutes=15))]
        perto = a[(a.t >= t - pd.Timedelta(hours=2)) & (a.t <= t + pd.Timedelta(minutes=15))]
        L.append(dict(
            t=t, duracao_h=round((fim - t).total_seconds() / 3600, 1),
            rodando_h=round((t - partida.index[-1]).total_seconds() / 3600, 1) if len(partida) else None,
            critico=bool(x.Prioridade.isin(["HIGH", "LOLO"]).any()),
            curado=next((f.strftime("%d/%m/%Y") for f in ft if abs((t - f).total_seconds()) <= 3 * 3600), ""),
            alarmes="; ".join(f"{r['Tag Alarme']} ({r['Descrição Alarme']}, {r.Prioridade})"
                              for _, r in perto.drop_duplicates("Tag Alarme").iterrows()),
            ocsvm_48h=bool(((eps.end >= t - pd.Timedelta(hours=48)) & (eps.start <= t)).any())))
    return pd.DataFrame(L)


def planilha(P: pd.DataFrame):
    rev, ref = P[~P.critico].reset_index(drop=True), P[P.critico].reset_index(drop=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Instruções"
    texto = [
        ("Paradas da turbina TC-33003A sem alarme crítico (07/2024 a 04/2026)", NEGRITO),
        ("", FONTE),
        (f"A aba 'Paradas para revisão' tem as {len(rev)} paradas em que a máquina saiu de operação sem alarme "
         "HIGH/LOLO do catálogo em ±15 min.", FONTE),
        ("Pedido: preencher as colunas amarelas (Classificação, Motivo e Fonte) de cada linha.", FONTE),
        ("Classificação: escolher na lista (Programada / Operacional (não programada, sem falha) / "
         "Falha de equipamento / Não sei).", FONTE),
        ("Fonte: de onde veio a informação (livro de turno, nota/ordem SAP, registro de partidas e paradas, "
         "memória da equipe...).", FONTE),
        ("Horários: 'Início (Brasília)' = UTC-3; 'Início (UTC)' é o relógio dos dados.", FONTE),
        ("Parada = estado operacional passa a 'desligada por longo período' depois de estar em operação "
         "nas 3 h anteriores (derivado dos sinais, não de registro da operação).", FONTE),
        ("A aba 'Trips curados (referência)' mostra as 8 paradas já reconhecidas como trip, só para comparação.",
         FONTE),
        ("", FONTE),
        ("Exemplo de preenchimento (fictício):", NEGRITO),
        ("Classificação = Programada · Motivo = Parada para lavagem do compressor · Fonte = Livro de turno "
         "de 12/03/2025", FONTE),
    ]
    for i, (s, f) in enumerate(texto, 1):
        ws.cell(i, 1, s).font = f
    ws.column_dimensions["A"].width = 130

    cab = ["Nº", "Início (Brasília)", "Início (UTC)", "Duração da parada (h)", "Em operação antes (h, desde a última parada longa)",
           "Alarmes do catálogo de 2 h antes a 15 min depois", "Classificação", "Motivo / observação", "Fonte"]
    larg = [5, 18, 18, 12, 12, 70, 30, 40, 25]

    def aba(nome, D, preencher):
        w = wb.create_sheet(nome)
        cols = cab if preencher else cab[:6] + ["Trip curado"]
        for j, (c, l) in enumerate(zip(cols, larg), 1):
            k = w.cell(1, j, c)
            k.font, k.alignment = NEGRITO, Alignment(wrap_text=True, vertical="top")
            w.column_dimensions[k.column_letter].width = l
        for i, r in D.iterrows():
            linha = [i + 1, (r.t - pd.Timedelta(hours=3)).to_pydatetime(), r.t.to_pydatetime(), r.duracao_h,
                     r.rodando_h, r.alarmes or "-"] + (["", "", ""] if preencher else [r.curado])
            for j, v in enumerate(linha, 1):
                k = w.cell(i + 2, j, v)
                k.font, k.alignment = FONTE, Alignment(wrap_text=True, vertical="top")
                if j in (2, 3):
                    k.number_format = "dd/mm/yyyy hh:mm"
                if preencher and j >= 7:
                    k.fill = AMARELO
        w.freeze_panes = "B2"
        if preencher:
            dv = DataValidation(type="list", formula1='"' + ",".join(OPCOES) + '"', allow_blank=True)
            w.add_data_validation(dv)
            dv.add(f"G2:G{len(D) + 1}")

    aba("Paradas para revisão", rev, True)
    aba("Trips curados (referência)", ref, False)
    wb.save(SAIDA / "paradas_para_revisao.xlsx")
    return rev, ref


def main():
    SAIDA.mkdir(exist_ok=True)
    P = paradas()
    P.to_csv(SAIDA / "paradas_interno.csv", index=False)
    rev, ref = planilha(P)
    print(f"paradas vindas de operação: {len(P)}; com alarme crítico: {len(ref)} "
          f"(batem com trip curado: {(ref.curado != '').sum()}); para revisão: {len(rev)}")
    print(f"OCSVM alertou nas 48 h antes: revisão {rev.ocsvm_48h.sum()}/{len(rev)}, trips {ref.ocsvm_48h.sum()}/{len(ref)}")


if __name__ == "__main__":
    main()
